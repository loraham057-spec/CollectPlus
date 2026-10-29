from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
STATIC_DIR = BASE_DIR / "static"
DATA_DIR = ROOT_DIR / "data"
DB_PATH = Path(os.getenv("COLLECTPLUS_DB", str(DATA_DIR / "collectplus.db")))
DATA_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(
    title="CollectPlus",
    version="0.1.0",
    description="Plateforme de collecte mobile et de visualisation des enquêtes.",
)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=False, allow_methods=["*"], allow_headers=["*"])
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


MAPEPI_FORM: dict[str, Any] = {
    "key": "mapepi-incidents-v1",
    "title": "Fiche d’enregistrement et de suivi d’incident",
    "short_title": "Incidents MAPEPI",
    "description": "Déclaration, suivi et clôture des incidents liés au transport et à la traçabilité des échantillons biologiques.",
    "version": 1,
    "sections": [
        {
            "title": "1. Identification de l’incident",
            "description": "Renseignez les informations disponibles au moment de la déclaration.",
            "questions": [
                {"name": "reference_incident", "label": "Référence de l’incident", "type": "text", "required": True, "placeholder": "Ex. INC-HL-2026-001"},
                {"name": "numero_epid", "label": "N° EPID", "type": "text", "required": False},
                {"name": "date_detection", "label": "Date de détection", "type": "date", "required": True},
                {"name": "heure_detection", "label": "Heure de détection", "type": "time"},
                {"name": "date_declaration", "label": "Date de déclaration", "type": "date", "required": True},
                {"name": "heure_declaration", "label": "Heure de déclaration", "type": "time"},
                {"name": "province", "label": "Province", "type": "select", "required": True, "options": ["Haut Lomami", "Tanganyika", "Kasaï", "Maniema", "Kinshasa", "Autre"]},
                {"name": "zone_sante", "label": "Zone de santé", "type": "text", "required": True},
                {"name": "site_prelevement", "label": "Site de prélèvement / aire de santé / structure", "type": "text", "required": True},
                {"name": "destination", "label": "Point de destination / laboratoire", "type": "text", "required": True},
                {"name": "declare_par", "label": "Déclaré par (nom et fonction)", "type": "text", "required": True},
                {"name": "telephone", "label": "Téléphone / contact", "type": "phone"},
                {"name": "organisation", "label": "Organisation / service", "type": "text"},
            ],
        },
        {
            "title": "2. Étape du circuit concernée",
            "description": "Plusieurs réponses sont possibles.",
            "questions": [
                {"name": "etapes_circuit", "label": "Étape(s) du circuit", "type": "multiselect", "required": True, "options": [
                    "Transport Communauté/AS vers le BCZ", "Transport BCZ vers l’antenne PEV", "Transport Communauté/AS vers l’antenne PEV", "Transport Antenne PEV vers Chef-Lieu Province ou point de sortie", "Transport Chef-Lieu Province vers Labo", "Réception au BCZ", "Réception à l’antenne PEV", "Réception au laboratoire d’analyse", "Autre"
                ]},
                {"name": "autre_etape", "label": "Autre étape (préciser)", "type": "text"},
            ],
        },
        {
            "title": "3. Type et niveau de l’incident",
            "description": "Sélectionnez tous les motifs constatés.",
            "questions": [
                {"name": "motifs_incident", "label": "Motif(s) de l’incident", "type": "multiselect", "required": True, "options": [
                    "Retard dans la collecte ou l’acheminement de l’échantillon", "Doublon du numéro EPID", "Échantillon mal étiqueté ou non étiqueté", "Échantillon prélevé dans un contenant inapproprié", "Réception d’un seul échantillon de selles au lieu des deux requis (PFA)", "Échantillon non localisé ou perte de traçabilité", "Fiche papier incomplète ou mal renseignée", "Échantillon reçu sans fiche d’accompagnement ou de notification", "Données DHIS2 manquantes ou erronées", "Absence de réseau ou difficulté de synchronisation des données", "Discordance entre la fiche papier et les données du Tracker DHIS2", "Problème de conditionnement ou d’emballage de l’échantillon", "Fuite ou déversement de l’échantillon", "Volume d’échantillon insuffisant", "Non-respect de la chaîne du froid pendant l’acheminement", "Accident de circulation impliquant le transporteur", "Refus de réception de l’échantillon par la structure destinataire", "Indisponibilité du moyen de transport", "Incident lié à la sécurité sur l’itinéraire", "Autre"
                ]},
                {"name": "autre_motif", "label": "Autre motif (préciser)", "type": "textarea"},
                {"name": "maladie_echantillon", "label": "Type de maladie / échantillon concerné", "type": "multiselect", "required": True, "options": ["PFA", "Cas contacts PFA", "Mpox", "Choléra", "Rougeole", "Fièvre jaune", "Fièvres hémorragiques virales"]},
            ],
        },
        {
            "title": "4. Impact sur le transport et la traçabilité",
            "description": "Cochez les conséquences observées ou les risques identifiés.",
            "questions": [
                {"name": "impacts", "label": "Impact(s) constaté(s)", "type": "multiselect", "options": [
                    "Aucun impact constaté", "Dépassement du délai recommandé de collecte ou d’acheminement", "Retard dans la transmission au laboratoire destinataire", "Risque de rejet de l’échantillon par le laboratoire", "Altération de la qualité ou de l’intégrité de l’échantillon", "Perte de traçabilité ou interruption du suivi", "Données de suivi incomplètes, manquantes ou erronées", "Exposition accidentelle du personnel à un échantillon biologique", "Risque de contamination du personnel ou de l’environnement", "Rupture de la chaîne du froid pendant l’acheminement", "Nécessité de procéder à un nouveau prélèvement", "Autre"
                ]},
                {"name": "autre_impact", "label": "Autre impact (préciser)", "type": "textarea"},
                {"name": "verification_effectuee", "label": "Vérification avec la fiche papier / le registre / le terrain effectuée ?", "type": "select", "required": True, "options": ["Oui", "Non"]},
            ],
        },
        {
            "title": "5. Notification et suivi",
            "questions": [
                {"name": "partage_whatsapp", "label": "Incident partagé au groupe WhatsApp de suivi ?", "type": "select", "options": ["Oui", "Non"]},
                {"name": "date_partage_whatsapp", "label": "Date du partage", "type": "date"},
                {"name": "presente_reunion", "label": "Présenté à la réunion hebdomadaire / bi-hebdomadaire ?", "type": "select", "options": ["Oui", "Non"]},
                {"name": "date_reunion", "label": "Date de la réunion", "type": "date"},
            ],
        },
        {
            "title": "6. Mesures immédiates prises",
            "description": "Saisissez les mesures principales. Le détail peut être complété dans les commentaires.",
            "questions": [
                {"name": "mesure_immediate", "label": "Mesure(s) immédiate(s) prise(s)", "type": "textarea", "required": True},
                {"name": "responsable_mesure", "label": "Responsable", "type": "text"},
                {"name": "date_heure_mesure", "label": "Date et heure", "type": "datetime"},
                {"name": "resultat_mesure", "label": "Résultat obtenu", "type": "textarea"},
            ],
        },
        {
            "title": "7. Mesures correctives et clôture",
            "questions": [
                {"name": "action_corrective", "label": "Action corrective / préventive convenue", "type": "textarea", "required": True},
                {"name": "responsable_action", "label": "Responsable de l’action", "type": "text", "required": True},
                {"name": "echeance_action", "label": "Échéance", "type": "date"},
                {"name": "statut_incident", "label": "Statut de l’incident", "type": "select", "required": True, "options": ["Ouvert", "En cours", "Résolu"]},
                {"name": "date_resolution", "label": "Date de résolution", "type": "date"},
                {"name": "preuve_cloture", "label": "Preuve / commentaire de clôture", "type": "textarea"},
            ],
        },
    ],
}


def db_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db() -> None:
    with db_conn() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS forms (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            short_title TEXT NOT NULL,
            description TEXT NOT NULL DEFAULT '',
            schema_json TEXT NOT NULL,
            created_at TEXT NOT NULL,
            is_system INTEGER NOT NULL DEFAULT 0
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS submissions (
            id TEXT PRIMARY KEY,
            form_id INTEGER NOT NULL,
            payload_json TEXT NOT NULL,
            submitted_at TEXT NOT NULL,
            FOREIGN KEY(form_id) REFERENCES forms(id)
        )""")
        row = conn.execute("SELECT id FROM forms WHERE short_title = ?", (MAPEPI_FORM["short_title"],)).fetchone()
        if row is None:
            conn.execute(
                "INSERT INTO forms(title, short_title, description, schema_json, created_at, is_system) VALUES (?, ?, ?, ?, ?, 1)",
                (MAPEPI_FORM["title"], MAPEPI_FORM["short_title"], MAPEPI_FORM["description"], json.dumps(MAPEPI_FORM, ensure_ascii=False), utc_now()),
            )


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def form_to_dict(row: sqlite3.Row) -> dict[str, Any]:
    schema = json.loads(row["schema_json"])
    schema["id"] = row["id"]
    schema["created_at"] = row["created_at"]
    schema["is_system"] = bool(row["is_system"])
    return schema


def get_form_or_404(form_id: int) -> dict[str, Any]:
    with db_conn() as conn:
        row = conn.execute("SELECT * FROM forms WHERE id = ?", (form_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Enquête introuvable")
    return form_to_dict(row)


class FormCreate(BaseModel):
    title: str = Field(min_length=2, max_length=160)
    short_title: str = Field(min_length=2, max_length=80)
    description: str = Field(default="", max_length=1000)
    sections: list[dict[str, Any]]


class SubmissionCreate(BaseModel):
    id: str | None = None
    form_id: int
    data: dict[str, Any]
    submitted_at: str | None = None


@app.on_event("startup")
def on_startup() -> None:
    init_db()


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/manifest.webmanifest")
def manifest() -> FileResponse:
    return FileResponse(STATIC_DIR / "manifest.webmanifest", media_type="application/manifest+json")


@app.get("/service-worker.js")
def service_worker() -> FileResponse:
    return FileResponse(STATIC_DIR / "service-worker.js", media_type="application/javascript")


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "app": "CollectPlus", "version": app.version, "time": utc_now()}


@app.get("/api/forms")
def list_forms() -> list[dict[str, Any]]:
    with db_conn() as conn:
        rows = conn.execute("SELECT * FROM forms ORDER BY id ASC").fetchall()
    return [form_to_dict(row) for row in rows]


@app.get("/api/forms/{form_id}")
def get_form(form_id: int) -> dict[str, Any]:
    return get_form_or_404(form_id)


@app.post("/api/forms")
def create_form(form: FormCreate) -> dict[str, Any]:
    if not form.sections:
        raise HTTPException(status_code=400, detail="Ajoutez au moins une section au formulaire.")
    for section in form.sections:
        if not isinstance(section, dict) or not section.get("questions"):
            raise HTTPException(status_code=400, detail="Chaque section doit contenir au moins une question.")
        for question in section["questions"]:
            if not question.get("name") or not question.get("label") or not question.get("type"):
                raise HTTPException(status_code=400, detail="Chaque question doit avoir un nom, un libellé et un type.")
    schema = {
        "key": f"custom-{uuid.uuid4().hex[:8]}",
        "title": form.title.strip(),
        "short_title": form.short_title.strip(),
        "description": form.description.strip(),
        "version": 1,
        "sections": form.sections,
    }
    with db_conn() as conn:
        cur = conn.execute(
            "INSERT INTO forms(title, short_title, description, schema_json, created_at, is_system) VALUES (?, ?, ?, ?, ?, 0)",
            (schema["title"], schema["short_title"], schema["description"], json.dumps(schema, ensure_ascii=False), utc_now()),
        )
        form_id = cur.lastrowid
    return get_form_or_404(int(form_id))


@app.delete("/api/forms/{form_id}")
def delete_form(form_id: int) -> dict[str, Any]:
    with db_conn() as conn:
        row = conn.execute("SELECT is_system FROM forms WHERE id = ?", (form_id,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Enquête introuvable")
        if row["is_system"]:
            raise HTTPException(status_code=400, detail="Le formulaire MAPEPI intégré ne peut pas être supprimé.")
        conn.execute("DELETE FROM submissions WHERE form_id = ?", (form_id,))
        conn.execute("DELETE FROM forms WHERE id = ?", (form_id,))
    return {"ok": True}


def parse_xlsform(content: bytes, filename: str) -> dict[str, Any]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Impossible de lire le fichier XLSForm : {exc}")
    if "survey" not in workbook.sheetnames:
        raise HTTPException(status_code=400, detail="Le fichier Excel doit contenir une feuille 'survey'.")
    survey_sheet = workbook["survey"]
    rows = list(survey_sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="La feuille survey est vide.")
    headers = [str(v).strip().lower() if v is not None else "" for v in rows[0]]
    table = [dict(zip(headers, row)) for row in rows[1:]]
    settings: dict[str, Any] = {}
    if "settings" in workbook.sheetnames:
        srows = list(workbook["settings"].iter_rows(values_only=True))
        if srows:
            sh = [str(v).strip().lower() if v is not None else "" for v in srows[0]]
            if len(srows) > 1:
                settings = dict(zip(sh, srows[1]))
    choices: dict[str, list[dict[str, str]]] = {}
    if "choices" in workbook.sheetnames:
        crows = list(workbook["choices"].iter_rows(values_only=True))
        if crows:
            ch = [str(v).strip().lower() if v is not None else "" for v in crows[0]]
            for crow in crows[1:]:
                rec = dict(zip(ch, crow))
                lname = str(rec.get("list_name") or "")
                if lname:
                    choices.setdefault(lname, []).append({"value": str(rec.get("name") or ""), "label": str(rec.get("label") or rec.get("name") or "")})

    sections: list[dict[str, Any]] = [{"title": "Questionnaire", "questions": []}]
    section_stack: list[dict[str, Any]] = [sections[0]]
    unsupported = []
    for record in table:
        raw_type = str(record.get("type") or "").strip()
        if not raw_type or raw_type.startswith("#"):
            continue
        name = str(record.get("name") or "").strip()
        label = str(record.get("label::French (fr)") or record.get("label::fr") or record.get("label") or name).strip()
        if raw_type == "begin_group":
            new_section = {"title": label or name or "Section", "description": "", "questions": []}
            sections.append(new_section)
            section_stack = [new_section]
            continue
        if raw_type == "end_group":
            section_stack = [sections[0]]
            continue
        if not name or not label:
            continue
        required_value = str(record.get("required") or "").strip().lower()
        required = required_value in ("yes", "true", "1")
        hint = str(record.get("hint::French (fr)") or record.get("hint") or "").strip()
        q: dict[str, Any] = {"name": name, "label": label, "required": required}
        if hint:
            q["hint"] = hint
        if raw_type.startswith("select_one "):
            q["type"] = "select"
            list_name = raw_type.split(" ", 1)[1].strip()
            q["options"] = [item["label"] for item in choices.get(list_name, [])]
        elif raw_type.startswith("select_multiple "):
            q["type"] = "multiselect"
            list_name = raw_type.split(" ", 1)[1].strip()
            q["options"] = [item["label"] for item in choices.get(list_name, [])]
        elif raw_type in ("integer", "decimal"):
            q["type"] = "number"
        elif raw_type in ("date", "time", "datetime", "text", "email", "note", "geopoint", "image", "audio", "barcode"):
            q["type"] = {"integer": "number", "decimal": "number", "note": "note", "geopoint": "text", "image": "text", "audio": "text", "barcode": "text"}.get(raw_type, raw_type)
        else:
            q["type"] = "text"
            unsupported.append(raw_type)
        if q["type"] in ("select", "multiselect") and not q.get("options"):
            q["options"] = ["Valeur 1", "Valeur 2"]
            q["hint"] = (q.get("hint", "") + " Liste de choix à vérifier après importation.").strip()
        section_stack[-1]["questions"].append(q)
    sections = [sec for sec in sections if sec.get("questions")]
    if not sections:
        raise HTTPException(status_code=400, detail="Aucune question exploitable n’a été trouvée dans la feuille survey.")
    title = str(settings.get("form_title") or settings.get("form_id") or Path(filename).stem).strip()
    return {
        "key": f"import-{uuid.uuid4().hex[:8]}",
        "title": title,
        "short_title": title[:80],
        "description": f"Formulaire importé depuis {filename}. Vérifiez les règles conditionnelles et les types de questions après importation.",
        "version": 1,
        "sections": sections,
        "import_warnings": sorted(set(unsupported)),
    }


@app.post("/api/forms/import")
async def import_form(file: UploadFile = File(...)) -> dict[str, Any]:
    filename = file.filename or "formulaire"
    content = await file.read()
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        try:
            schema = json.loads(content.decode("utf-8-sig"))
        except Exception:
            raise HTTPException(status_code=400, detail="Le JSON n’est pas valide.")
        if not isinstance(schema, dict) or not isinstance(schema.get("sections"), list):
            raise HTTPException(status_code=400, detail="Le JSON doit contenir un objet avec un tableau 'sections'.")
        schema.setdefault("title", Path(filename).stem)
        schema.setdefault("short_title", str(schema["title"])[:80])
        schema.setdefault("description", "Formulaire importé depuis JSON.")
    elif suffix in (".xlsx", ".xlsm"):
        schema = parse_xlsform(content, filename)
    else:
        raise HTTPException(status_code=400, detail="Format non pris en charge. Utilisez un XLSForm .xlsx ou un schéma .json CollectPlus.")
    form = FormCreate(title=str(schema.get("title", "Enquête importée")), short_title=str(schema.get("short_title", schema.get("title", "Enquête importée")))[:80], description=str(schema.get("description", "")), sections=schema["sections"])
    created = create_form(form)
    if schema.get("import_warnings"):
        created["import_warnings"] = schema["import_warnings"]
    return created


@app.post("/api/submissions")
def create_submission(submission: SubmissionCreate) -> dict[str, Any]:
    get_form_or_404(submission.form_id)
    sub_id = submission.id or str(uuid.uuid4())
    submitted_at = submission.submitted_at or utc_now()
    payload = submission.data
    with db_conn() as conn:
        try:
            conn.execute(
                "INSERT INTO submissions(id, form_id, payload_json, submitted_at) VALUES (?, ?, ?, ?)",
                (sub_id, submission.form_id, json.dumps(payload, ensure_ascii=False), submitted_at),
            )
        except sqlite3.IntegrityError:
            # Retry-safe offline synchronization: an already uploaded UUID is an acknowledgement.
            row = conn.execute("SELECT id FROM submissions WHERE id = ?", (sub_id,)).fetchone()
            if row is None:
                raise HTTPException(status_code=409, detail="Impossible d’enregistrer cette soumission.")
    return {"ok": True, "id": sub_id, "submitted_at": submitted_at}


def read_submissions(form_id: int | None = None) -> list[dict[str, Any]]:
    with db_conn() as conn:
        if form_id:
            rows = conn.execute("""SELECT s.*, f.short_title, f.title AS form_title FROM submissions s
                                    JOIN forms f ON f.id=s.form_id WHERE s.form_id=? ORDER BY s.submitted_at DESC""", (form_id,)).fetchall()
        else:
            rows = conn.execute("""SELECT s.*, f.short_title, f.title AS form_title FROM submissions s
                                    JOIN forms f ON f.id=s.form_id ORDER BY s.submitted_at DESC""").fetchall()
    return [{"id": r["id"], "form_id": r["form_id"], "form_title": r["form_title"], "short_title": r["short_title"], "data": json.loads(r["payload_json"]), "submitted_at": r["submitted_at"]} for r in rows]


@app.get("/api/submissions")
def list_submissions(form_id: int | None = None) -> list[dict[str, Any]]:
    return read_submissions(form_id)


@app.get("/api/dashboard")
def dashboard() -> dict[str, Any]:
    with db_conn() as conn:
        form_count = conn.execute("SELECT COUNT(*) FROM forms").fetchone()[0]
        submission_count = conn.execute("SELECT COUNT(*) FROM submissions").fetchone()[0]
        recent = conn.execute("""SELECT s.id, s.form_id, s.payload_json, s.submitted_at, f.short_title
                                FROM submissions s JOIN forms f ON f.id=s.form_id ORDER BY s.submitted_at DESC LIMIT 7""").fetchall()
        by_form = conn.execute("""SELECT f.short_title, COUNT(s.id) AS n FROM forms f LEFT JOIN submissions s ON s.form_id=f.id
                               GROUP BY f.id ORDER BY n DESC, f.id ASC""").fetchall()
        all_payloads = conn.execute("SELECT payload_json FROM submissions").fetchall()
    statuses: dict[str, int] = {"Ouvert": 0, "En cours": 0, "Résolu": 0}
    disease_counts: dict[str, int] = {}
    for row in all_payloads:
        payload = json.loads(row["payload_json"])
        status = payload.get("statut_incident")
        if status in statuses:
            statuses[status] += 1
        diseases = payload.get("maladie_echantillon", [])
        if isinstance(diseases, str):
            diseases = [diseases]
        if isinstance(diseases, list):
            for disease in diseases:
                disease_counts[str(disease)] = disease_counts.get(str(disease), 0) + 1
    return {
        "form_count": form_count,
        "submission_count": submission_count,
        "status_counts": statuses,
        "disease_counts": disease_counts,
        "by_form": [{"label": r["short_title"], "value": r["n"]} for r in by_form],
        "recent": [{"id": r["id"], "form_id": r["form_id"], "form_title": r["short_title"], "data": json.loads(r["payload_json"]), "submitted_at": r["submitted_at"]} for r in recent],
    }


def flatten_records(rows: list[dict[str, Any]]) -> tuple[list[str], list[list[str]]]:
    all_keys: list[str] = []
    for row in rows:
        for key in row["data"].keys():
            if key not in all_keys:
                all_keys.append(key)
    headers = ["id", "form_id", "form_title", "submitted_at"] + all_keys
    output = []
    for row in rows:
        values = [row["id"], row["form_id"], row["form_title"], row["submitted_at"]]
        for key in all_keys:
            value = row["data"].get(key, "")
            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            values.append("" if value is None else str(value))
        output.append(values)
    return headers, output


@app.get("/api/export/csv")
def export_csv(form_id: int | None = None) -> StreamingResponse:
    rows = read_submissions(form_id)
    headers, values = flatten_records(rows)
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(values)
    # UTF-8 BOM helps Excel open French accents correctly.
    content = "\ufeff" + buffer.getvalue()
    name = f"collectplus_export_{form_id or 'toutes_enquetes'}.csv"
    return StreamingResponse(iter([content]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/export/json")
def export_json(form_id: int | None = None) -> StreamingResponse:
    rows = read_submissions(form_id)
    content = json.dumps(rows, ensure_ascii=False, indent=2)
    name = f"collectplus_export_{form_id or 'toutes_enquetes'}.json"
    return StreamingResponse(iter([content]), media_type="application/json; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/export/xlsx")
def export_xlsx(form_id: int | None = None) -> StreamingResponse:
    rows = read_submissions(form_id)
    headers, values = flatten_records(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Soumissions"
    sheet.append(headers)
    for row in values:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        max_len = max((len(str(cell.value or "")) for cell in column), default=10)
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
    out = io.BytesIO()
    workbook.save(out)
    out.seek(0)
    name = f"collectplus_export_{form_id or 'toutes_enquetes'}.xlsx"
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{name}"'})
